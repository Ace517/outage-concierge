#!/usr/bin/env python3
"""build_data.py - turn the raw AEP shapefiles into the two small tables the app reads.

You only need this if you want to rebuild data/processed/. The processed files are already included.

    pip install geopandas pyogrio openpyxl
    python build_data.py --raw path/to/aep_challenge_hack26/data/raw

Privacy: premise numbers are replaced by a random hash (the salt is never saved), and the
life-support flag, house numbers and coordinates never leave this script.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import json
import secrets
import tempfile
import warnings
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

import engine

warnings.filterwarnings("ignore")

STORM_START, STORM_END = pd.Timestamp("2025-02-12"), pd.Timestamp("2025-02-19")  # end exclusive
LAST_REAL_REPAIR = pd.Timestamp("2025-02-20")  # later dates in the file look like open-ticket placeholders
OUT = Path(__file__).parent / "data" / "processed"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", default="../aep_challenge_hack26/data/raw", help="folder with Data1.zip, Data2.zip, matched_upstream.zip")
    ap.add_argument("--workbook", default=None, help="optional path to Install_Labor_HackOHIO.xlsx to cross-check labor constants")
    args = ap.parse_args()

    cust, storm, cc, prem, sys_dpk = load(Path(args.raw))
    write(storm, cc, prem, sys_dpk)
    if args.workbook:
        verify_labor(args.workbook)


def load(raw: Path):
    """Read the raw zips and return (customers with geometry, storm damage, circuits table, premises table, system damage rate).

    The premises table has one row per customer row, in the same order. build_map.py relies on that.
    """
    import geopandas as gpd

    tmp = Path(tempfile.mkdtemp(prefix="aep_"))
    for z, sub in [("Data1.zip", "d1"), ("Data2.zip", "d2"), ("matched_upstream.zip", "up")]:
        with zipfile.ZipFile(raw / z) as f:
            f.extractall(tmp / sub)
    print("extracted to", tmp)

    # ---------------- load ----------------
    cust = gpd.read_file(tmp / "d1" / "Customer_Clip.shp")
    dmg = gpd.read_file(tmp / "d1" / "arcos_prod_damage_09_28_Clip.shp")

    cust["circuit"] = pd.to_numeric(cust["circuit_id"], errors="coerce")
    cust = cust.dropna(subset=["circuit"]).copy()
    cust["circuit"] = cust["circuit"].astype(int)
    cust["premise"] = cust["premise_nu"].astype(str)

    storm = dmg[(dmg.Report_Tim >= STORM_START) & (dmg.Report_Tim < STORM_END)].copy()
    storm["circuit"] = pd.to_numeric(storm.Circuit.str.extract(r"(\d+)")[0], errors="coerce")
    storm = storm.dropna(subset=["circuit"]).copy()
    storm["circuit"] = storm["circuit"].astype(int)
    storm["cat"] = storm.Asset_Type.map(engine.ASSET_TO_CAT).fillna("Wire span")
    storm["access"] = (storm.Access_Iss == "Yes").astype(float)
    storm["special"] = storm.Special_Eq.notna().astype(float)
    print("storm damage records:", len(storm), "| circuits:", storm.circuit.nunique())

    # ---------------- upstream topology ----------------
    chain = {}
    dev2prem = collections.defaultdict(set)
    with open(tmp / "up" / "matched_upstream.txt", encoding="utf-8", errors="ignore") as f:
        for line in f:
            p = line.rstrip("\r\n").split("\t")
            items = [x.split("*") for x in p[1:] if x]
            chain[p[0]] = items
            for x in items:
                if len(x) >= 3:
                    dev2prem[x[1]].add(p[0])
    protective = lambda t: t.startswith("Fuse") or t.startswith("Recloser") or t.startswith("Sectionalizer")
    devs = {pr: sum(1 for x in it if len(x) >= 3 and protective(x[2])) for pr, it in chain.items()}

    # ---------------- which premises sit downstream of a device that was damaged? ----------------
    layers = {"Fuse Bank": "d1/Fuse_Bank_Clip.shp", "Transformer": "d2/Transformer_Clip.shp",
              "Recloser Bank": "d2/Recloser_Bank_Clip.shp", "Switch": "d2/Switch_Clip.shp"}
    hit_prem = set()
    for asset, rel in layers.items():
        G = gpd.read_file(tmp / rel).to_crs(3857)[["id", "geometry"]].rename(columns={"id": "dev_id"})
        S = storm[storm.Asset_Type == asset].to_crs(3857)[["geometry"]]
        if not len(S):
            continue
        j = gpd.sjoin_nearest(S, G, how="left", distance_col="m", max_distance=30)
        for d in j[j.m.notna()].dev_id.dropna().astype("int64").astype(str).unique():
            hit_prem |= dev2prem.get(d, set())
    print("premises downstream of a damaged device:", len(hit_prem))

    # ---------------- circuits table ----------------
    cc = cust.groupby("circuit").size().rename("customers").to_frame()
    sys_mix = storm.cat.value_counts(normalize=True).reindex(engine.CATS).fillna(0.0)
    sys_access, sys_special = storm.access.mean(), storm.special.mean()
    n_by = storm.groupby("circuit").size()
    cc["dmg"] = n_by.reindex(cc.index).fillna(0).astype(int)
    cc["dpk"] = cc.dmg / cc.customers * 1000
    cc["pct"] = cc.dpk.rank(pct=True)
    sys_dpk = storm.shape[0] / cc.customers.sum() * 1000
    cc["base_items"] = np.maximum(1.0, 0.7 * cc.dmg + 0.3 * sys_dpk * cc.customers / 1000)
    counts = storm.groupby(["circuit", "cat"]).size().unstack(fill_value=0).reindex(columns=engine.CATS, fill_value=0)
    counts = counts.reindex(cc.index, fill_value=0)
    K = 3.0  # shrink small circuits toward the system-wide mix
    for c in engine.CATS:
        cc["mix_" + c] = (counts[c] + K * sys_mix[c]) / (cc.dmg + K)
    acc = storm.groupby("circuit").access.sum().reindex(cc.index).fillna(0)
    spe = storm.groupby("circuit").special.sum().reindex(cc.index).fillna(0)
    cc["access_share"] = (acc + 5 * sys_access) / (cc.dmg + 5)
    cc["special_share"] = (spe + 5 * sys_special) / (cc.dmg + 5)

    # ---------------- premises table ----------------
    salt = secrets.token_hex(16)  # random, never saved: hashes cannot be reversed from premise numbers
    for n in (8, 10, 12):
        hid = cust.premise.map(lambda s: hashlib.sha256((salt + s).encode()).hexdigest()[:n].upper())
        if hid.is_unique:
            break
    prem = pd.DataFrame({
        "hid": hid.values,
        "city": cust.city.fillna("??").str.title().replace({"??": "Town not recorded"}).values,
        "road": cust.road_name.fillna("Unnamed road").str.title().replace({"??": "Unnamed Road"}).values,
        "circuit": cust.circuit.values,
        "svc": cust.kind_of_se.fillna("Unknown").values,
        "devs": cust.premise.map(lambda s: devs.get(s, 0)).values,
        "chain_ok": cust.premise.isin(chain.keys()).values,
        "hit25": cust.premise.isin(hit_prem).values,
    })
    prem["pct"] = prem.circuit.map(cc.pct)
    return cust, storm, cc, prem, sys_dpk


def write(storm, cc, prem, sys_dpk):
    # ---------------- history benchmarks and back-test ----------------
    rep = storm[(storm.Repaired == "Yes") & (storm.Repaired_D <= LAST_REAL_REPAIR)].copy()
    rep["days"] = (rep.Repaired_D - rep.Report_Tim).dt.days
    q = rep.days.quantile([0.5, 0.9, 0.99]).to_dict()

    span = rep.groupby("circuit").apply(lambda g: (g.Repaired_D.max() - g.Report_Tim.min()).days)
    hits, rows = 0, 0
    for c, d_obs in span.items():
        if c not in cc.index:
            continue
        r = cc.loc[c]
        n_items = max(1, int(r.dmg))
        mix = {k: r["mix_" + k] for k in engine.CATS}
        e = engine.restore(n_items, mix, r.access_share, r.special_share, engine.default_crews(n_items), engine.OBSERVED_SEV)
        obs_lo, obs_hi = d_obs * 24, d_obs * 24 + 24
        rows += 1
        hits += int(e["lo_h"] <= obs_hi and e["hi_h"] >= obs_lo)
    backtest = {"circuits": rows, "overlap": hits, "share": round(hits / rows, 3) if rows else None}
    print("backtest:", backtest)

    # ---------------- personas (fictional names, real-data-backed profiles) ----------------
    def E_of(row):
        return engine.exposure(row.pct, row.hit25, row.svc, row.devs, row.chain_ok)[2]

    prem["E"] = prem.apply(E_of, axis=1)
    ok = prem[prem.chain_ok]
    a = ok[(ok.svc == "Overhead") & ok.hit25 & (ok.pct >= 0.9)].sort_values(["E", "hid"], ascending=[False, True]).iloc[0]
    b = ok[(ok.svc == "Underground") & ~ok.hit25 & (ok.pct <= 0.25)].sort_values(["E", "hid"]).iloc[0]
    mid = ok[(ok.svc == "Overhead") & ~ok.hit25 & (ok.pct.between(0.35, 0.65))].copy()
    mid["gap"] = (mid.E - 0.45).abs()
    c = mid.sort_values(["gap", "hid"]).iloc[0]
    hit = ok[(ok.svc == "Underground") & ok.hit25 & (ok.pct.between(0.35, 0.65))].copy()
    hit["gap"] = (hit.pct - 0.5).abs()
    d = hit.sort_values(["gap", "hid"]).iloc[0]
    personas = [
        {"name": "Maria", "hid": a.hid, "medical": True,
         "blurb": "Someone at home uses medical equipment. Overhead service on one of the hardest-hit circuits."},
        {"name": "Devon", "hid": b.hid, "medical": False,
         "blurb": "Underground service on a circuit that held up well in the last big storm."},
        {"name": "Priya", "hid": c.hid, "medical": False,
         "blurb": "Overhead service on a mid-range circuit."},
        {"name": "David", "hid": d.hid, "medical": False, "enrolled": True,
         "blurb": "Underground service, but equipment that feeds his home was damaged in the last big storm. Already signed up for alerts."},
    ]

    meta = {
        "premises": int(len(prem)), "circuits": int(len(cc)),
        "storm_records": int(len(storm)), "storm_circuits_hit": int((cc.dmg > 0).sum()),
        "storm_window": "Feb 12 to 18, 2025",
        "direct_hit_premises": int(prem.hit25.sum()),
        "direct_hit_share": round(float(prem.hit25.mean()), 4),
        "circuit_exposed_share": round(float((prem.circuit.map(cc.dmg) > 0).mean()), 3),
        "underground_share": round(float((prem.svc == "Underground").mean()), 3),
        "repair_days_p50": float(q[0.5]), "repair_days_p90": float(q[0.9]), "repair_days_p99": float(q[0.99]),
        "sys_dpk": round(float(sys_dpk), 2),
        "backtest": backtest, "personas": personas,
    }

    OUT.mkdir(parents=True, exist_ok=True)
    prem.drop(columns=["E", "pct"]).to_csv(OUT / "premises.csv.gz", index=False)
    cc.reset_index().round(5).to_csv(OUT / "circuits.csv", index=False)
    (OUT / "meta.json").write_text(json.dumps(meta, indent=2))
    print(json.dumps({k: v for k, v in meta.items() if k != "personas"}, indent=2))
    print("personas:", [(p["name"], p["hid"]) for p in personas])


def verify_labor(path):
    """Confirm the labor constants in engine.py appear in the supplied workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    vals = []
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            vals += [float(v) for v in row if isinstance(v, (int, float))]
    missing = [(k, x) for k, pair in engine.WORKBOOK_HOURS.items() for x in pair
               if not any(abs(x - v) < 1e-6 for v in vals)]
    print("labor constants all found in workbook" if not missing else f"NOT FOUND in workbook: {missing}")


if __name__ == "__main__":
    main()
